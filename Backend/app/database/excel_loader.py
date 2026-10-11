from __future__ import annotations

import re
import math
import tempfile
import uuid
import unicodedata
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy import Boolean, Column, DateTime, Float, Integer, MetaData, String, Table, create_engine, insert

from app.database.database_manager import database_manager


MAX_FILE_SIZE = 25 * 1024 * 1024
MAX_PREVIEW_ROWS = 8
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm"}


def _slug(value: str, fallback: str) -> str:
    value = re.sub(r"[^0-9A-Za-zÀ-ÿ_]+", "_", str(value).strip()).strip("_").lower()
    if value and value[0].isdigit():
        value = f"col_{value}"
    return value or fallback


def _json_value(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _infer_type(values: list[Any]) -> str:
    clean = [value for value in values if value not in (None, "")]
    if not clean:
        return "empty"
    if all(isinstance(value, bool) for value in clean):
        return "boolean"
    if all(isinstance(value, int) and not isinstance(value, bool) for value in clean):
        return "integer"
    if all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in clean):
        return "number"
    if all(isinstance(value, (datetime, date)) for value in clean):
        return "date"
    return "text"


def _looks_like_header(row: list[Any]) -> bool:
    filled = [v for v in row if v not in (None, "")]
    if not filled:
        return False
    return len(filled) >= 2 and all(isinstance(v, str) and bool(v.strip()) for v in filled)


def _normalize_sheet(sheet) -> dict:
    rows = list(sheet.iter_rows(values_only=True))
    non_empty = [row for row in rows if any(value not in (None, "") for value in row)]
    if not non_empty:
        return {"name": sheet.title, "table_name": _slug(sheet.title, "sheet"), "headers": [], "data_rows": [], "schema": [], "warnings": ["La hoja está vacía."]}

    header_row = list(non_empty[0])
    has_header = _looks_like_header(header_row)
    if not has_header:
        header_row = [f'column_{i + 1}' for i in range(max(len(row) for row in non_empty))]
    width = max(len(row) for row in non_empty)
    headers: list[str] = []
    used: set[str] = set()
    warnings: list[str] = []
    if not has_header:
        warnings.append('No se identificó una fila de encabezados; se generaron nombres de columna y se conservaron todas las filas.')

    for index in range(width):
        raw = header_row[index] if index < len(header_row) else None
        base = _slug(raw, f"column_{index + 1}") if raw not in (None, "") else f"column_{index + 1}"
        name = base
        suffix = 2
        while name in used:
            name = f"{base}_{suffix}"
            suffix += 1
        if raw in (None, ""):
            warnings.append(f"La columna {index + 1} no tiene encabezado; se interpretó como {name}.")
        used.add(name)
        headers.append(name)

    source_rows = non_empty[1:] if has_header else non_empty
    data_rows = [list(row) + [None] * (width - len(row)) for row in source_rows]
    schema = []
    for index, header in enumerate(headers):
        values = [row[index] for row in data_rows]
        non_null = sum(value not in (None, "") for value in values)
        schema.append({
            "name": header,
            "source_header": _json_value(header_row[index]) if index < len(header_row) else None,
            "type": _infer_type(values),
            "non_null": non_null,
            "nulls": len(values) - non_null,
            "unique_values": len({str(value) for value in values if value not in (None, "")}),
            "mixed_types": len({type(value).__name__ for value in values if value not in (None, "")}) > 1,
        })
    if any(col["mixed_types"] and col["type"] == "text" for col in schema):
        warnings.append("Hay columnas con tipos mezclados; se conservarán como texto para evitar pérdida de información.")
    return {"name": sheet.title, "table_name": _slug(sheet.title, "sheet"), "headers": headers, "data_rows": data_rows, "schema": schema, "warnings": warnings}



def _semantic_name(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", str(value).lower())
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "_", normalized).strip("_")


def _semantic_profile(sheet: dict) -> dict:
    """Sugerencias explicables; nunca cambian los datos ni asignan roles definitivos."""
    from app.analysis.semantic_mapper_v2 import ROLE_PROFILES

    table_name = _semantic_name(sheet["name"])
    column_names = {_semantic_name(col["name"]) for col in sheet["schema"]}
    candidates = []
    for domain, roles in ROLE_PROFILES.items():
        for role, profile in roles.items():
            names = {_semantic_name(term) for term in profile.table_terms}
            matched_table = table_name in names or any(
                table_name.startswith(term + "_") or table_name.endswith("_" + term)
                for term in names if len(term) >= 4
            )
            matched_columns = sorted(
                col for col in column_names
                if any(col == term or col.startswith(term + "_") or col.endswith("_" + term)
                       for term in profile.column_terms if len(term) >= 3)
            )
            score = (4 if matched_table else 0) + min(3, len(matched_columns))
            if score >= 3:
                candidates.append({
                    "domain": domain, "role": role,
                    "confidence": "high" if score >= 5 else "medium",
                    "evidence": {"sheet_name": matched_table, "columns": matched_columns[:6]},
                    "_score": score,
                })
    candidates.sort(key=lambda item: (-item["_score"], item["domain"], item["role"]))
    return {"suggestions": [{key: value for key, value in item.items() if key != "_score"}
                            for item in candidates[:5]], "requires_confirmation": True}


def _suggest_relations(sheets: list[dict]) -> list[dict]:
    """Propone relaciones por nombres y valores, sin crear claves foráneas."""
    suggestions = []
    for child in sheets:
        for column in child["schema"]:
            key = _semantic_name(column["name"])
            if not key.endswith("_id") or key == "id":
                continue
            prefix = key[:-3]
            child_values = {str(row[child["headers"].index(column["name"])])
                            for row in child["data_rows"] if row[child["headers"].index(column["name"])] not in (None, "")}
            if not child_values:
                continue
            for parent in sheets:
                if parent is child:
                    continue
                parent_name = _semantic_name(parent["name"])
                if not (parent_name == prefix or parent_name == prefix + "s" or
                        parent_name.rstrip("es").rstrip("s") == prefix.rstrip("es").rstrip("s")):
                    continue
                for parent_key in ("id", key):
                    if parent_key not in parent["headers"]:
                        continue
                    index = parent["headers"].index(parent_key)
                    values = [str(row[index]) for row in parent["data_rows"] if row[index] not in (None, "")]
                    if not values or len(values) != len(set(values)):
                        continue
                    coverage = len(child_values.intersection(values)) / len(child_values)
                    if coverage >= 0.8:
                        suggestions.append({
                            "from_sheet": child["name"], "from_column": column["name"],
                            "to_sheet": parent["name"], "to_column": parent_key,
                            "match_ratio": round(coverage, 3), "confirmed": False,
                        })
                        break
    return suggestions[:30]


def _public_profile(sheet: dict) -> dict:
    preview = [{sheet["headers"][i]: _json_value(row[i]) for i in range(len(sheet["headers"]))} for row in sheet["data_rows"][:MAX_PREVIEW_ROWS]]
    return {"name": sheet["name"], "table_name": sheet["table_name"], "rows": len(sheet["data_rows"]), "columns": len(sheet["headers"]), "headers": sheet["headers"], "schema": sheet["schema"], "preview": preview, "warnings": sheet["warnings"], "semantic": _semantic_profile(sheet)}


async def _read_workbook(file: UploadFile) -> tuple[str, bytes, list[dict]]:
    filename = file.filename or "workbook.xlsx"
    if Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ValueError("Formato no compatible. Usa un archivo .xlsx o .xlsm.")
    content = await file.read(MAX_FILE_SIZE + 1)
    if not content:
        raise ValueError("El archivo Excel está vacío.")
    if len(content) > MAX_FILE_SIZE:
        raise ValueError("El archivo Excel supera el límite de 25 MB.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheets = [_normalize_sheet(sheet) for sheet in workbook.worksheets]
        workbook.close()
    except Exception as exc:
        raise ValueError(f"No se pudo leer el libro de Excel: {exc}") from exc
    return filename, content, sheets


async def inspect_excel_upload(file: UploadFile) -> dict:
    filename, content, sheets = await _read_workbook(file)
    profiles = [_public_profile(sheet) for sheet in sheets]
    populated = [sheet for sheet in profiles if sheet["rows"] > 0]
    total_rows = sum(sheet["rows"] for sheet in populated)
    return {
        "uploaded": True, "source_type": "excel", "filename": filename, "size_bytes": len(content),
        "sheet_count": len(profiles), "populated_sheets": len(populated), "total_data_rows": total_rows,
        "total_columns": sum(sheet["columns"] for sheet in populated), "sheets": profiles,
        "suggested_relations": _suggest_relations(sheets),
        "interpretation": {"summary": f"Se detectaron {len(populated)} hojas con datos y {total_rows} filas.", "ready_to_import": bool(populated)},
    }


def _sql_type(kind: str):
    return {"boolean": Boolean, "integer": Integer, "number": Float, "date": DateTime}.get(kind, String)


def _db_value(value: Any, kind: str) -> Any:
    if value in (None, ""):
        return None
    if kind == "date":
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, time.min)
        return str(value)
    if kind == "text" or kind == "empty":
        return str(value)
    return value


async def activate_excel_upload(file: UploadFile) -> dict:
    """Convierte un Excel en SQLite local y lo activa como la fuente SQL del sistema."""
    filename, content, sheets = await _read_workbook(file)
    populated = [sheet for sheet in sheets if sheet["data_rows"] and sheet["headers"]]
    if not populated:
        raise ValueError("El libro no contiene hojas con filas de datos utilizables.")

    temp_dir = Path(tempfile.gettempdir()) / "kenneth_excel_sources"
    temp_dir.mkdir(parents=True, exist_ok=True)
    safe_name = _slug(Path(filename).stem, "excel")
    db_path = temp_dir / f"{safe_name}_{uuid.uuid4().hex[:12]}.sqlite3"

    engine = create_engine(f"sqlite:///{db_path}")
    metadata = MetaData()
    tables: dict[str, Table] = {}
    used_tables: set[str] = set()

    for sheet_index, sheet in enumerate(populated, start=1):
        base = sheet["table_name"] or f"sheet_{sheet_index}"
        table_name = base
        suffix = 2
        while table_name in used_tables:
            table_name = f"{base}_{suffix}"
            suffix += 1
        used_tables.add(table_name)
        sheet["table_name"] = table_name
        columns = [Column(item["name"], _sql_type(item["type"]), nullable=True) for item in sheet["schema"]]
        tables[table_name] = Table(table_name, metadata, *columns)

    metadata.create_all(engine)
    try:
        with engine.begin() as connection:
            for sheet in populated:
                table = tables[sheet["table_name"]]
                batch = []
                for row in sheet["data_rows"]:
                    batch.append({item["name"]: _db_value(row[index], item["type"]) for index, item in enumerate(sheet["schema"])})
                    if len(batch) >= 500:
                        connection.execute(insert(table), batch)
                        batch.clear()
                if batch:
                    connection.execute(insert(table), batch)
    except Exception:
        engine.dispose()
        db_path.unlink(missing_ok=True)
        raise

    total_rows = sum(len(sheet["data_rows"]) for sheet in populated)
    status = database_manager.configure_excel(engine=engine, filename=filename, tables=list(tables), rows=total_rows, database_path=str(db_path))
    return {
        "uploaded": True,
        "activated": True,
        "source_type": "excel",
        "connection": status,
        "filename": filename,
        "size_bytes": len(content),
        "tables": [_public_profile(sheet) for sheet in populated],
        "suggested_relations": _suggest_relations(populated),
        "total_data_rows": total_rows,
    }
